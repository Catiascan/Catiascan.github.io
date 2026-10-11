"""Converte Budget_FSA_Hub.xlsx nos documentos do app Gestao FSA (db do artifact).

Uso:
    python3 -I tools/gestao_fsa/convert.py <Budget_FSA_Hub.xlsx> <pasta_saida>

Gera um ficheiro JSON por documento em <pasta_saida>/<colecao>/<id>.json e um
manifest.json com a lista de writes, em lotes de no maximo 50 (limite do
ArtifactData batch). Nenhum dado financeiro fica neste repositorio: a planilha
e a saida ficam fora do git.
"""
import datetime as dt
import json
import re
import sys
import unicodedata
from pathlib import Path

import openpyxl

MESES = ["jan", "fev", "mar", "abr", "mai", "jun", "jul", "ago", "set", "out", "nov", "dez"]


def norm(s):
    s = unicodedata.normalize("NFKD", str(s or "")).encode("ascii", "ignore").decode()
    return s.lower().strip()


def iso(v):
    """Datetime ou 'dd/mm/aaaa' -> 'aaaa-mm-dd'; o resto fica None."""
    if isinstance(v, (dt.datetime, dt.date)):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, str):
        m = re.fullmatch(r"\s*(\d{1,2})/(\d{1,2})/(\d{4})\s*", v)
        if m:
            d, mo, y = m.groups()
            return f"{y}-{int(mo):02d}-{int(d):02d}"
    return None


def num(v):
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return round(float(v), 2)
    return None


def txt(v):
    if v is None:
        return ""
    if isinstance(v, (dt.datetime, dt.date)):
        return v.strftime("%d/%m/%Y")
    return str(v).replace("Â·", "·").strip()


# --- Plano de contas -------------------------------------------------------
# Regras por palavra-chave na descricao, aplicadas quando a categoria original
# e generica ("Saida"). A ordem importa: a primeira que casa vence.
REGRAS_DESC = [
    (r"transferencia (enviada|recebida)|transf\. interna", "Transferência", ""),
    (r"rendimento", "Rendimento/Aplicação", ""),
    (r"imovel|aluguel|financiamento", "Custo fixo", "Imóvel"),
    (r"edp|energia", "Custo fixo", "Energia"),
    (r"claro", "Custo fixo", "Internet/telefone"),
    (r"saae|agua", "Custo fixo", "Água"),
    (r"porto seguro saude|convenio medico", "Custo fixo", "Convênio médico"),
    (r"seg vida|portoseg|seguro", "Custo fixo", "Seguros"),
    (r"tarifa|iof|juros", "Bancário", ""),
    (r"protesto|cartorio", "Tributos", "Passivo/protestos"),
    (r"kitano", "Comissões", "Kitano"),
    (r"sonia", "Empréstimos", "Sônia"),
    (r"matheus", "Custo variável", "Freelancer"),
    (r"posto|estacionamento", "Transporte", "Combustível/estacionamento"),
    (r"daise", "Empréstimos", "Daise"),
    (r"fernando sampaio|catia scandola", "Adiantamento a sócio", ""),
    (r"ifood|bar|lanchonete|padaria|pontual|meep|farmacia|drogaria|hematologia|"
     r"cipax|pernambucanas|c e a|belian|dondoka|recreacao", "Pessoal (PF)", ""),
]

# Categoria original (normalizada) -> (grupo, subcategoria)
def grupo_por_categoria(cat, tag):
    c = norm(cat)
    t = (tag or "").upper()
    if t == "TRANSFERENCIA" or c.startswith("transferencia"):
        return "Transferência", ""
    if t == "ESTORNO" or c.startswith("estorno"):
        return "Estorno", ""
    if t in ("RENDIMENTO", "APLICACAO") or c in ("rendimento", "aplicacao financeira"):
        return "Rendimento/Aplicação", ""
    if t in ("ENTRADA_CLIENTE", "ENTRADA") or c == "entrada":
        return "Receita", ""
    if t in ("ENTRADA_OUTRA", "ENTRADA_PF") or c == "entrada pf":
        return "Outras entradas", ""
    if "adiantamento" in c or "acerto socio" in c or t == "ADIANTAMENTO":
        return "Adiantamento a sócio", ""
    if "pro-labore" in c:
        return "Pró-labore", ""
    if "tribut" in c or c.startswith("passivo"):
        return "Tributos", ""
    if c.startswith("custo fixo"):
        m = re.search(r"\(([^)]*)\)", c)
        return "Custo fixo", (m.group(1).split(",")[0].strip().capitalize() if m else "")
    if c.startswith("custo variavel") or c == "custo operacional":
        return "Custo variável", "Operação"
    if c.startswith("transporte"):
        return "Transporte", "Uber/99"
    if c == "pessoal":
        return "Pessoal (PF)", ""
    if c == "comissao":
        return "Comissões", ""
    if c.startswith("emprestimo"):
        return "Empréstimos", ""
    return None, ""


def classificar(desc, cat, tag):
    g, sub = grupo_por_categoria(cat, tag)
    if g:
        return g, sub
    d = norm(desc)
    for padrao, grupo, s in REGRAS_DESC:
        if re.search(padrao, d):
            return grupo, s
    return "A classificar", ""


def slug(s):
    return re.sub(r"[^A-Za-z0-9_-]+", "_", norm(s)).strip("_")[:60] or "x"


# --- Conversores por aba ---------------------------------------------------

def lancamentos(ws):
    out = {}
    for i, r in enumerate(ws.iter_rows(min_row=3, values_only=True), start=3):
        r = list(r) + [None] * 10
        data = iso(r[0])
        valor = num(r[5])
        if not data or valor is None:
            continue
        grupo, sub = classificar(r[2], r[3], r[9])
        out[f"rd-{i:04d}"] = {
            "data": data, "mes": data[:7], "banco": txt(r[1]), "descricao": txt(r[2]),
            "categoria_original": txt(r[3]), "grupo": grupo, "sub": sub,
            "modo": txt(r[4]), "valor": valor, "cliente": txt(r[8]), "tag": txt(r[9]),
            "origem": f"Registro_Diario!{i}",
        }
    return out


def contas_pagar(ws):
    out = {}
    for i, r in enumerate(ws.iter_rows(min_row=4, values_only=True), start=4):
        r = list(r) + [None] * 10
        venc = iso(r[0])
        if not venc:
            continue
        out[f"cp-{i:04d}"] = {
            "vencimento": venc, "mes": venc[:7], "descricao": txt(r[1]), "categoria": txt(r[2]),
            "previsto": num(r[3]) or 0, "conta": txt(r[4]), "baixa": iso(r[5]) or "",
            "pago": num(r[6]), "estado": txt(r[7]).upper() or "PENDENTE",
            "no_registro": txt(r[8]), "obs": txt(r[9]), "origem": f"Contas_a_Pagar!{i}",
        }
    return out


def contas_receber(ws):
    out = {}
    for i, r in enumerate(ws.iter_rows(min_row=4, max_row=12, values_only=True), start=4):
        r = list(r) + [None] * 8
        dev = txt(r[0])
        if not dev or dev.upper().startswith("TOTAL") or num(r[1]) is None:
            continue
        out[f"cr-{i:04d}"] = {
            "devedor": dev, "total": num(r[1]) or 0, "ja_pago": num(r[2]) or 0,
            "aberto": num(r[3]) or 0, "previsto": iso(r[4]) or "", "previsto_txt": txt(r[4]),
            "obs": txt(r[5]), "baixa": iso(r[6]) or "", "estado": txt(r[7]).upper() or "EM ABERTO",
            "origem": f"Contas_a_Receber!{i}",
        }
    return out


def budget(ws):
    """So linhas de lancamento; totais e resultado o app calcula."""
    out, secao, ordem = {}, "", 0
    for i, r in enumerate(ws.iter_rows(min_row=4, values_only=True), start=4):
        item = txt(r[0])
        if not item or item.startswith("Item ("):
            continue
        if re.match(r"^\d\.", item):
            secao = item
            continue
        if item.startswith("=") or secao.startswith(("5.", "6.")):
            continue
        meses = [num(v) or 0 for v in r[1:13]]
        ordem += 1
        out[f"b-{ordem:02d}"] = {
            "ano": 2026, "secao": secao, "item": item, "ordem": ordem,
            "meses": meses, "origem": f"Budget_2026!{i}",
        }
    return out


def clientes(ws):
    out = {}
    for i, r in enumerate(ws.iter_rows(min_row=3, values_only=True), start=3):
        nome = txt(r[0])
        if not nome or nome in ("Projeto",) or nome.upper() == "TOTAL" or nome.lower().startswith("saldo"):
            if nome.upper() == "TOTAL":
                break
            continue
        status = txt(r[10])
        # Ativo = ainda tem receita prevista de outubro em diante.
        ativo = any((num(v) or 0) > 0 for v in r[6:9])
        out[f"cli-{slug(nome)}"] = {
            "nome": nome, "status": "Ativo" if ativo else "Encerrado",
            "previsto": {m: num(v) or 0 for m, v in zip(MESES[4:], r[1:9])},
            "obs": status, "origem": f"Receitas_Esperadas!{i}",
        }
    return out


def passivo(ws):
    out = {}
    for i, r in enumerate(ws.iter_rows(min_row=4, values_only=True), start=4):
        nome = txt(r[0])
        if not nome or nome.upper().startswith(("TOTAL", "CONTESTADA", "NOTA", "ATUALIZADO")):
            continue
        if num(r[2]) is None:
            continue
        out[f"pv-{i:02d}"] = {
            "divida": nome, "direcao": txt(r[1]), "original": num(r[2]) or 0,
            "negociavel": num(r[3]) or 0, "fase": txt(r[4]), "status": txt(r[5]),
            "obs": txt(r[6]), "origem": f"Passivo!{i}",
        }
    return out


def abas(wb, limite=180_000):
    """Copia fiel (valores) de todas as abas, para consulta. Divide abas grandes."""
    out = {}
    for pos, ws in enumerate(wb.worksheets):
        linhas = []
        for r in ws.iter_rows(values_only=True):
            cel = []
            for v in r:
                n = num(v)
                cel.append(n if n is not None else txt(v))
            while cel and cel[-1] == "":
                cel.pop()
            linhas.append({"c": cel})
        while linhas and not linhas[-1]["c"]:
            linhas.pop()
        partes, atual, tam = [], [], 0
        for ln in linhas:
            t = len(json.dumps(ln, ensure_ascii=False).encode())
            if atual and tam + t > limite:
                partes.append(atual)
                atual, tam = [], 0
            atual.append(ln)
            tam += t
        partes.append(atual)
        base = slug(ws.title)
        for k, p in enumerate(partes):
            out[base if k == 0 else f"{base}~{k + 1}"] = {
                "nome": ws.title, "posicao": pos, "parte": k + 1, "partes": len(partes),
                "linha_inicial": 1 + sum(len(x) for x in partes[:k]), "linhas": p,
            }
    return out


def main(xlsx, saida):
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    m = wb["Mes_Atual"]
    colecoes = {
        "lancamentos": lancamentos(wb["Registro_Diario"]),
        "pagar": contas_pagar(wb["Contas_a_Pagar"]),
        "receber": contas_receber(wb["Contas_a_Receber"]),
        "budget": budget(wb["Budget_2026"]),
        "clientes": clientes(wb["Receitas_Esperadas"]),
        "passivo": passivo(wb["Passivo"]),
        "abas": abas(wb),
        "config": {"geral": {
            "saldo_pj": num(m["C6"].value) or 0,
            "saldo_obs": txt(m["D6"].value)[:600],
            "saldo_piso": 15000,
            "fonte": Path(xlsx).name,
            "importado_em": dt.date.today().isoformat(),
        }},
    }
    saida = Path(saida)
    writes = []
    for col, docs in colecoes.items():
        (saida / col).mkdir(parents=True, exist_ok=True)
        for doc_id, data in docs.items():
            f = saida / col / f"{doc_id}.json"
            f.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
            writes.append({"op": "set", "collection": col, "doc_id": doc_id, "file_path": str(f)})
    lotes = [writes[i:i + 50] for i in range(0, len(writes), 50)]
    (saida / "manifest.json").write_text(json.dumps(lotes, ensure_ascii=False, indent=0), encoding="utf-8")
    for col, docs in colecoes.items():
        print(f"{col}: {len(docs)}")
    print(f"total: {len(writes)} documentos em {len(lotes)} lotes")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
